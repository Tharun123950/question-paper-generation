from docx import Document
from flask import send_file, request, jsonify
import os
import random
import hashlib
import uuid
import re
import pdfplumber
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
from datetime import datetime
from flask import Flask
from flask_sqlalchemy import SQLAlchemy
from flask_cors import CORS
import jwt


app = Flask(__name__)

CORS(app)

app.config["SQLALCHEMY_DATABASE_URI"] = "postgresql://postgres:prasanna123@localhost:5432/question_generator"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["SECRET_KEY"] = "your-secret-key-change-in-production"

db = SQLAlchemy(app)

# ------------------ MODELS ------------------
class User(db.Model):
    __tablename__ = "users"
    
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password = db.Column(db.String(255), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    
    def set_password(self, password):
        self.password = generate_password_hash(password)
    
    def check_password(self, password):
        return check_password_hash(self.password, password)


class GeneratedPaper(db.Model):

    id = db.Column(db.Integer, primary_key=True)

    fingerprint = db.Column(db.String(200))
    signature_hash = db.Column(db.String(200))
    paper_content = db.Column(db.Text)

    year = db.Column(db.String(20))
    education = db.Column(db.String(50))
    semester = db.Column(db.String(50))
    exam_name = db.Column(db.String(100))
    month_year = db.Column(db.String(50))

    subject_name = db.Column(db.String(100))

    exam_time = db.Column(db.String(50))
    max_marks = db.Column(db.String(20))
    branch_section = db.Column(db.String(50))
    exam_date = db.Column(db.String(30))

    regulation = db.Column(db.String(20))

    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class Question(db.Model):
    __tablename__ = "questions"

    id = db.Column(db.Integer, primary_key=True)
    unit = db.Column(db.Integer, nullable=False)
    marks = db.Column(db.Integer, nullable=False)
    question_text = db.Column(db.Text, nullable=False)

# ------------------ HELPER FUNCTIONS ------------------

def generate_signature(question_ids):
    sorted_ids = sorted(question_ids)
    signature_string = "-".join(map(str, sorted_ids))
    return hashlib.sha256(signature_string.encode()).hexdigest()


def pick_random(question_list, count):
    if len(question_list) < count:
        return None
    return random.sample(question_list, count)


# ------------------ ROUTES ------------------

@app.route("/")
def home():
    return "PostgreSQL Connected Successfully ✅"


# ================= AUTHENTICATION ROUTES =================
@app.route("/api/register", methods=["POST"])
def register():
    try:
        data = request.get_json()
        
        if not data or not data.get("username") or not data.get("email") or not data.get("password"):
            return jsonify({"error": "Missing required fields"}), 400
        
        if User.query.filter_by(username=data["username"]).first():
            return jsonify({"error": "Username already exists"}), 400
        
        if User.query.filter_by(email=data["email"]).first():
            return jsonify({"error": "Email already exists"}), 400
        
        user = User(username=data["username"], email=data["email"])
        user.set_password(data["password"])
        
        db.session.add(user)
        db.session.commit()
        
        token = jwt.encode({
            "user_id": user.id,
            "username": user.username,
            "email": user.email
        }, app.config["SECRET_KEY"], algorithm="HS256")
        
        return jsonify({
            "message": "User registered successfully",
            "token": token,
            "user": {
                "id": user.id,
                "username": user.username,
                "email": user.email
            }
        }), 201
    except Exception as e:
        db.session.rollback()
        return jsonify({"error": str(e)}), 500


@app.route("/api/login", methods=["POST"])
def login():
    try:
        data = request.get_json()
        
        if not data or not data.get("username") or not data.get("password"):
            return jsonify({"error": "Missing username or password"}), 400
        
        user = User.query.filter_by(username=data["username"]).first()
        
        if not user or not user.check_password(data["password"]):
            return jsonify({"error": "Invalid username or password"}), 401
        
        token = jwt.encode({
            "user_id": user.id,
            "username": user.username,
            "email": user.email
        }, app.config["SECRET_KEY"], algorithm="HS256")
        
        return jsonify({
            "message": "Login successful",
            "token": token,
            "user": {
                "id": user.id,
                "username": user.username,
                "email": user.email
            }
        }), 200
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/logout", methods=["POST"])
def logout():
    # Frontend should clear token from localStorage
    return jsonify({"message": "Logout successful"}), 200


UPLOAD_FOLDER = "uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

from docx import Document  # add at top if not added

@app.route("/upload-pdf", methods=["POST"])
def upload_pdf():

    # ---------------- UNIT VALIDATION ----------------
    unit = request.args.get("unit")

    if not unit:
        return {"error": "Unit parameter missing. Use ?unit=1, 2, 3, 4, or 5"}, 400

    try:
        unit = int(unit)
    except ValueError:
        return {"error": "Unit must be an integer (1 to 5)"}, 400

    if unit not in [1, 2, 3, 4, 5]:
        return {"error": "Unit must be between 1 and 5"}, 400


    # ---------------- FILE VALIDATION ----------------
    if "file" not in request.files:
        return {"error": "No file uploaded"}, 400

    file = request.files["file"]

    if file.filename == "":
        return {"error": "Empty filename"}, 400

    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config["UPLOAD_FOLDER"], filename)
    file.save(filepath)

    current_marks = None
    questions_added = 0
    duplicates_skipped = 0


    # -------- HELPER FUNCTIONS (NEW) --------
    def normalize_text(text):
        text = text.lower()
        text = re.sub(r'[^a-z0-9\s]', '', text)
        text = re.sub(r'\s+', ' ', text).strip()
        return text

    def check_and_add_question(q_text, marks, u):
        if not q_text or marks is None:
            return 0, 0
            
        q_text = q_text.replace("\n", " ").strip()
        if not q_text:
            return 0, 0
            
        normalized_new = normalize_text(q_text)
        
        existing = Question.query.filter_by(unit=u, marks=marks).all()
        duplicate_found = False
        for q in existing:
            if normalize_text(q.question_text) == normalized_new:
                duplicate_found = True
                break
                
        if duplicate_found:
            return 0, 1 # added, skipped
            
        new_question = Question(
            unit=u,
            marks=marks,
            question_text=q_text
        )
        db.session.add(new_question)
        return 1, 0

    def get_marks_from_heading(text):
        text_upper = text.upper()
        if re.search(r'(SECTION|PART|SECTIOIN)\s*[-_:]?\s*[A]', text_upper):
            return 2
        elif re.search(r'(SECTION|PART|SECTIOIN)\s*[-_:]?\s*[B]', text_upper):
            return 5
        return None

    def extract_numbered_question(text):
        match = re.match(r'^\s*(?:[Qq]\s*\d+|\d+)\s*[\.\):-]\s*(.+)', text)
        if match:
            return match.group(1).strip()
        return None


    # =====================================================
    # ====================== PDF PARSING ==================
    # =====================================================
    if filename.lower().endswith(".pdf"):

        with pdfplumber.open(filepath) as pdf:
            current_question = None

            for page in pdf.pages:

                # 1. Plain Text Parsing (Line by Line)
                text = page.extract_text()
                if text:
                    lines = text.split('\n')
                    for line in lines:
                        # Check header
                        marks_candidate = get_marks_from_heading(line)
                        if marks_candidate:
                            # Flush before changing
                            if current_question and current_marks is not None:
                                added, skipped = check_and_add_question(current_question, current_marks, unit)
                                questions_added += added
                                duplicates_skipped += skipped
                            current_marks = marks_candidate
                            current_question = None
                            continue
                            
                        # Check question
                        q_text = extract_numbered_question(line)
                        if q_text is not None:
                            if current_question and current_marks is not None:
                                added, skipped = check_and_add_question(current_question, current_marks, unit)
                                questions_added += added
                                duplicates_skipped += skipped
                            current_question = q_text
                        elif current_question is not None:
                            # Ignore short random numbers like page numbers
                            if not re.match(r'^\s*\d+\s*$', line):
                                current_question += " " + line.strip()

                # 2. Table Parsing (Fallback for existing formats)
                tables = page.extract_tables()
                for table in tables:
                    for row in table:
                        if not row or len(row) < 2:
                            continue

                        sno = row[0]
                        question_text = row[1]

                        if sno is None or question_text is None:
                            continue

                        if str(sno).strip().lower() in ["sno", "no"]:
                            continue

                        if current_marks is not None:
                            added, skipped = check_and_add_question(question_text, current_marks, unit)
                            questions_added += added
                            duplicates_skipped += skipped

            # Flush last text question in PDF
            if current_question and current_marks is not None:
                added, skipped = check_and_add_question(current_question, current_marks, unit)
                questions_added += added
                duplicates_skipped += skipped


    # =====================================================
    # ===================== DOCX PARSING ==================
    # =====================================================
    elif filename.lower().endswith(".docx"):

        doc = Document(filepath)
        current_question = None

        # 1. Detect SECTION & Plain Text Tracking
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            marks_candidate = get_marks_from_heading(text)
            if marks_candidate:
                if current_question and current_marks is not None:
                    added, skipped = check_and_add_question(current_question, current_marks, unit)
                    questions_added += added
                    duplicates_skipped += skipped
                current_marks = marks_candidate
                current_question = None
                continue

            q_text = extract_numbered_question(text)
            if q_text is not None:
                if current_question and current_marks is not None:
                    added, skipped = check_and_add_question(current_question, current_marks, unit)
                    questions_added += added
                    duplicates_skipped += skipped
                current_question = q_text
            elif current_question is not None:
                current_question += " " + text
                
        # Flush last text question in DOCX
        if current_question and current_marks is not None:
            added, skipped = check_and_add_question(current_question, current_marks, unit)
            questions_added += added
            duplicates_skipped += skipped

        # 2. Parse tables like PDF
        for table in doc.tables:
            for row in table.rows:
                cells = row.cells
                if len(cells) < 2:
                    continue

                sno = cells[0].text.strip()
                question_text = cells[1].text.strip()

                if not sno or not question_text:
                    continue

                if sno.lower() in ["sno", "no"]:
                    continue

                if current_marks is not None:
                    added, skipped = check_and_add_question(question_text, current_marks, unit)
                    questions_added += added
                    duplicates_skipped += skipped

    else:
        return {"error": "Only PDF and DOCX files are supported"}, 400

    # ---------------- COMMIT ----------------
    db.session.commit()

    return {
        "message": "File processed successfully",
        "file_type": "PDF" if filename.lower().endswith(".pdf") else "DOCX",
        "unit_uploaded": unit,
        "questions_added": questions_added,
        "duplicates_skipped": duplicates_skipped
    }


@app.route("/generate-type1")
def generate_type1():

    # ================= HEADER PARAMETERS =================
    year = request.args.get("year")
    education = request.args.get("education")
    semester = request.args.get("semester")
    exam_name = request.args.get("exam_name")
    month_year = request.args.get("month_year")

    subject_name = request.args.get("subject_name")

    exam_time = request.args.get("exam_time")
    max_marks = request.args.get("max_marks")
    branch_section = request.args.get("branch_section")
    exam_date = request.args.get("exam_date")

    regulation = request.args.get("regulation")

    # ================= FETCH QUESTIONS =================
    unit1_2m = Question.query.filter_by(unit=1, marks=2).all()
    unit2_2m = Question.query.filter_by(unit=2, marks=2).all()
    unit1_5m = Question.query.filter_by(unit=1, marks=5).all()
    unit2_5m = Question.query.filter_by(unit=2, marks=5).all()

    # ================= VALIDATION =================
    if len(unit1_2m) < 2:
        return "Not enough Unit 1 short questions"
    if len(unit2_2m) < 2:
        return "Not enough Unit 2 short questions"
    if len(unit1_5m) < 6:
        return "Not enough Unit 1 long questions"
    if len(unit2_5m) < 6:
        return "Not enough Unit 2 long questions"

    # ---------------- PART A ---------------- 
    partA = [] 
    
    partA_u1 = pick_random(unit1_2m, 2) 
    partA.extend(partA_u1) 
    
    partA_u2 = pick_random(unit2_2m, 2) 
    partA.extend(partA_u2) 
    
    remaining_short = [q for q in (unit1_2m + unit2_2m) if q not in partA] 
    partA_mixed = pick_random(remaining_short, 1) 
    partA.extend(partA_mixed)

    # ================= PART B =================
    u1_pool = unit1_5m.copy() 
    u2_pool = unit2_5m.copy() 

    block1 = pick_random(u1_pool, 4) 
    q2 = block1[:2] 
    q3 = block1[2:] 
    u1_pool = [q for q in u1_pool if q not in block1] 
    
    block2 = pick_random(u2_pool, 4) 
    q4 = block2[:2] 
    q5 = block2[2:] 
    u2_pool = [q for q in u2_pool if q not in block2] 
    
    q6a = pick_random(u1_pool, 1) 
    u1_pool = [q for q in u1_pool if q not in q6a] 
    
    q6b = pick_random(u2_pool, 1) 
    u2_pool = [q for q in u2_pool if q not in q6b] 
    
    q7a = pick_random(u1_pool, 1) 
    u1_pool = [q for q in u1_pool if q not in q7a] 
    
    q7b = pick_random(u2_pool, 1) 
    u2_pool = [q for q in u2_pool if q not in q7b]

    # ================= FORMAT =================
    COLUMN_WIDTH = 90

    def format_line(label, question, marks):
        text = f"{label}) {question}"
        return text.ljust(COLUMN_WIDTH) + f"({marks}M)"

    output = ""

    output += "PART-A (2 Marks Each) - Compulsory\n\n"
    labelsA = ["1i", "1ii", "1iii", "1iv", "1v"]

    for i in range(5):
        output += format_line(labelsA[i], partA[i].question_text, 2) + "\n\n"

    output += "\nPART-B (5 Marks Each)\n\n"

    output += format_line("2a", q2[0].question_text, 5) + "\n\n"
    output += format_line("2b", q2[1].question_text, 5) + "\n\n"
    output += "OR\n\n"
    output += format_line("3a", q3[0].question_text, 5) + "\n\n"
    output += format_line("3b", q3[1].question_text, 5) + "\n\n"

    output += format_line("4a", q4[0].question_text, 5) + "\n\n"
    output += format_line("4b", q4[1].question_text, 5) + "\n\n"
    output += "OR\n\n"
    output += format_line("5a", q5[0].question_text, 5) + "\n\n"
    output += format_line("5b", q5[1].question_text, 5) + "\n\n"

    output += format_line("6a", q6a[0].question_text, 5) + "\n\n"
    output += format_line("6b", q6b[0].question_text, 5) + "\n\n"
    output += "OR\n\n"
    output += format_line("7a", q7a[0].question_text, 5) + "\n\n"
    output += format_line("7b", q7b[0].question_text, 5) + "\n\n"

    # ================= DUPLICATE CHECK =================
    selected = partA + q2 + q3 + q4 + q5 + q6a + q6b + q7a + q7b
    selected_ids = [q.id for q in selected]
    signature_hash = generate_signature(selected_ids)

    existing = GeneratedPaper.query.filter_by(signature_hash=signature_hash).first()
    if existing:
        return "Duplicate Paper Detected ❌"

    new_paper = GeneratedPaper(
        fingerprint=",".join(map(str, sorted(selected_ids))),
        signature_hash=signature_hash,
        paper_content=output,

        year=year,
        education=education,
        semester=semester,
        exam_name=exam_name,
        month_year=month_year,
        subject_name=subject_name,
        exam_time=exam_time,
        max_marks=max_marks,
        branch_section=branch_section,
        exam_date=exam_date,
        regulation=regulation
    )

    db.session.add(new_paper)
    db.session.commit()

    return {
        "paper_id": new_paper.id,
        "content": output
    }

@app.route("/generate-type2")
def generate_type2():

    # ================= HEADER PARAMETERS =================
    year = request.args.get("year")
    education = request.args.get("education")
    semester = request.args.get("semester")
    exam_name = request.args.get("exam_name")
    month_year = request.args.get("month_year")

    subject_name = request.args.get("subject_name")

    exam_time = request.args.get("exam_time")
    max_marks = request.args.get("max_marks")
    branch_section = request.args.get("branch_section")
    exam_date = request.args.get("exam_date")

    regulation = request.args.get("regulation")

    # ================= FETCH QUESTIONS =================
    unit3_2m = Question.query.filter_by(unit=3, marks=2).all()
    unit4_2m = Question.query.filter_by(unit=4, marks=2).all()
    unit5_2m = Question.query.filter_by(unit=5, marks=2).all()

    unit3_5m = Question.query.filter_by(unit=3, marks=5).all()
    unit4_5m = Question.query.filter_by(unit=4, marks=5).all()
    unit5_5m = Question.query.filter_by(unit=5, marks=5).all()

    # ================= VALIDATION =================
    if len(unit3_2m) < 2:
        return "Not enough Unit 3 short questions"
    if len(unit4_2m) < 2:
        return "Not enough Unit 4 short questions"
    if len(unit5_2m) < 1:
        return "Not enough Unit 5 short questions"
    
    if len(unit3_5m) < 4:
        return "Not enough Unit 3 long questions"
    if len(unit4_5m) < 4:
        return "Not enough Unit 4 long questions"
    if len(unit5_5m) < 4:
        return "Not enough Unit 5 long questions"

    # ---------------- PART A ---------------- 
    partA = [] 
    
    partA_u3 = pick_random(unit3_2m, 2) 
    partA.extend(partA_u3) 
    
    partA_u4 = pick_random(unit4_2m, 2) 
    partA.extend(partA_u4) 
    
    partA_u5 = pick_random(unit5_2m, 1) 
    partA.extend(partA_u5)

    # ================= PART B =================
    u3_pool = unit3_5m.copy() 
    u4_pool = unit4_5m.copy() 
    u5_pool = unit5_5m.copy() 

    block1 = pick_random(u3_pool, 4) 
    q2 = block1[:2] 
    q3 = block1[2:] 
    
    block2 = pick_random(u4_pool, 4) 
    q4 = block2[:2] 
    q5 = block2[2:] 
    
    block3 = pick_random(u5_pool, 4)
    q6 = block3[:2]
    q7 = block3[2:]

    # ================= FORMAT =================
    COLUMN_WIDTH = 90

    def format_line(label, question, marks):
        text = f"{label}) {question}"
        return text.ljust(COLUMN_WIDTH) + f"({marks}M)"

    output = ""

    output += "PART-A (2 Marks Each) - Compulsory\n\n"
    labelsA = ["1i", "1ii", "1iii", "1iv", "1v"]

    for i in range(5):
        output += format_line(labelsA[i], partA[i].question_text, 2) + "\n\n"

    output += "\nPART-B (5 Marks Each)\n\n"

    output += format_line("2a", q2[0].question_text, 5) + "\n\n"
    output += format_line("2b", q2[1].question_text, 5) + "\n\n"
    output += "OR\n\n"
    output += format_line("3a", q3[0].question_text, 5) + "\n\n"
    output += format_line("3b", q3[1].question_text, 5) + "\n\n"

    output += format_line("4a", q4[0].question_text, 5) + "\n\n"
    output += format_line("4b", q4[1].question_text, 5) + "\n\n"
    output += "OR\n\n"
    output += format_line("5a", q5[0].question_text, 5) + "\n\n"
    output += format_line("5b", q5[1].question_text, 5) + "\n\n"

    output += format_line("6a", q6[0].question_text, 5) + "\n\n"
    output += format_line("6b", q6[1].question_text, 5) + "\n\n"
    output += "OR\n\n"
    output += format_line("7a", q7[0].question_text, 5) + "\n\n"
    output += format_line("7b", q7[1].question_text, 5) + "\n\n"

    # ================= DUPLICATE CHECK =================
    selected = partA + q2 + q3 + q4 + q5 + q6 + q7
    selected_ids = [q.id for q in selected]
    signature_hash = generate_signature(selected_ids)

    existing = GeneratedPaper.query.filter_by(signature_hash=signature_hash).first()
    if existing:
        return "Duplicate Paper Detected ❌"

    new_paper = GeneratedPaper(
        fingerprint=",".join(map(str, sorted(selected_ids))),
        signature_hash=signature_hash,
        paper_content=output,

        year=year,
        education=education,
        semester=semester,
        exam_name=exam_name,
        month_year=month_year,
        subject_name=subject_name,
        exam_time=exam_time,
        max_marks=max_marks,
        branch_section=branch_section,
        exam_date=exam_date,
        regulation=regulation
    )

    db.session.add(new_paper)
    db.session.commit()

    return {
        "paper_id": new_paper.id,
        "content": output
    }

@app.route("/papers")
def get_papers():
    papers = GeneratedPaper.query.order_by(GeneratedPaper.created_at.desc()).all()

    result = []
    for paper in papers:
        result.append({
            "id": paper.id,
            "created_at": paper.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            "signature": paper.signature_hash
        })

    return {"papers": result}

@app.route("/download-paper/<int:paper_id>")
def download_paper(paper_id):

    paper = GeneratedPaper.query.get(paper_id)
    if not paper:
        return "Paper not found", 404

    format_type = request.args.get("format", "docx").lower()

    # ================= HEADER PARAMETERS =================
    year = paper.year
    education = paper.education
    semester = paper.semester
    exam_name = paper.exam_name
    month_year = paper.month_year

    subject_name = paper.subject_name

    exam_time = paper.exam_time
    max_marks = paper.max_marks
    branch_section = paper.branch_section
    exam_date = paper.exam_date

    regulation = paper.regulation

    content = paper.paper_content.strip()

    # Remove old text-based header from stored content
    if "PART-A" in content:
        content = content.split("PART-A", 1)[1]
        content = "PART-A" + content.strip()

    # =====================================================
    # ====================== PDF ==========================
    # =====================================================
    if format_type == "pdf":

        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph
        )
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.enums import TA_CENTER, TA_RIGHT,TA_LEFT
        from io import BytesIO
        import os
        import re

        buffer = BytesIO()

        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=36,
            leftMargin=36,
            topMargin=30,
            bottomMargin=30
        )

        elements = []
        styles = getSampleStyleSheet()

        # -------- LOGO + REGULATION --------
        from reportlab.platypus import Image, Table, TableStyle, Spacer
        from reportlab.lib.units import inch
        from reportlab.lib import colors

        logo = Image("static/logo.png")
        logo.drawWidth = 3.4 * inch
        logo.drawHeight = 0.6 * inch
        logo.hAlign = "CENTER"

        r23 = Table(
                [[Paragraph(f"<b>{regulation}</b>", styles["Normal"])]],
                colWidths=[0.6 * inch],
                rowHeights=[0.35 * inch]
        )

        r23.setStyle(TableStyle([
                ("BOX",(0,0),(-1,-1),1,colors.black),
                ("ALIGN",(0,0),(-1,-1),"CENTER"),
                ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ]))

        header_table = Table(
                [[logo, r23]],
                colWidths=[6.8 * inch, 0.6 * inch]
        )

        header_table.setStyle(TableStyle([
                ("VALIGN",(0,0),(-1,-1),"TOP"),
                ("ALIGN",(0,0),(0,0),"CENTER"),
                ("ALIGN",(1,0),(1,0),"RIGHT"),
                ("RIGHTPADDING",(1,0),(1,0),2),
                ("LEFTPADDING",(0,0),(-1,-1),0),
                ("TOPPADDING",(0,0),(-1,-1),0),
                ("BOTTOMPADDING",(0,0),(-1,-1),0),
        ]))

        elements.append(header_table)
        elements.append(Spacer(1,6))
        
                # -------- STREAM / EXAM TITLE --------
        stream_style = ParagraphStyle(
            name="StreamStyle",
            parent=styles["Normal"],
            fontSize=12,
            alignment=TA_CENTER,
            leading=14,
            spaceAfter=4
        )

        elements.append(
            Paragraph(
                f"<b>{year} {education}; {semester}; {exam_name}; {month_year}</b>",
                stream_style
            )
        )
        elements.append(Spacer(1,10))

        # -------- SUBJECT --------
        subject_style = ParagraphStyle(
            name="SubjectStyle",
            parent=styles["Normal"],
            fontSize=14,
            alignment=TA_CENTER,
            leading=16,
            spaceAfter=6
        )

        elements.append(
            Paragraph(
                f"<b>{subject_name.upper()}</b>",
                subject_style
            )
        )

        elements.append(Spacer(1,10))

        # -------- TIME / MARKS TABLE --------

        right_align_style = ParagraphStyle(
            name="RightAlign",
            parent=styles["Normal"],
            alignment=TA_RIGHT
        )

        left_align_style = ParagraphStyle(
            name="LeftAlign",
            parent=styles["Normal"],
            alignment=TA_LEFT
        )

        info_data = [
            [
                Paragraph(f"<b>Time: {exam_time}</b>", left_align_style),
                Paragraph(f"<b>Max. Marks: {max_marks}</b>", right_align_style),
            ],
            [
                Paragraph(f"<b>Branch & Section: {branch_section}</b>", left_align_style),
                Paragraph(f"<b>Date: {exam_date}</b>", right_align_style),
            ],
        ]

        info_table = Table(
            info_data,
            colWidths=[3.5 * inch, 3.5 * inch]
        )

        info_table.setStyle(TableStyle([
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),

            ("LEFTPADDING", (0,0), (-1,-1), 0),
            ("RIGHTPADDING", (0,0), (-1,-1), 0),

            ("TOPPADDING", (0,0), (-1,-1), 2),
            ("BOTTOMPADDING", (0,0), (-1,-1), 2),
        ]))

        elements.append(info_table)
        elements.append(Spacer(1,6))
        # -------- REGISTER NUMBER --------
        reg_boxes = [["" for _ in range(8)]]

        boxes_table = Table(reg_boxes, colWidths=20, rowHeights=18)
        boxes_table.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 1, colors.black),
            ("INNERGRID", (0, 0), (-1, -1), 1, colors.black),
        ]))

        register_table = Table(
            [[Paragraph("<b>Registered No</b>", styles["Normal"]), boxes_table]],
            colWidths=[1.8 * inch, 5.2 * inch]
        )

        register_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))

        elements.append(register_table)
        elements.append(Spacer(1, 4))

        # -------- SEPARATOR LINE --------
        separator = Table([[""]], colWidths=[520])
        separator.setStyle(TableStyle([
            ("LINEBELOW", (0, 0), (-1, -1), 1, colors.black)
        ]))

        elements.append(separator)
        elements.append(Spacer(1, 8))

        # -------- INSTRUCTIONS --------
        instruction_style = ParagraphStyle(
            name="InstructionStyle",
            parent=styles["Normal"],
            alignment=TA_CENTER,
            fontSize=11,
            leading=14,
            spaceAfter=6
        )

        elements.append(Paragraph(
            "<b>Question Paper Consists of Part-A and Part-B</b>",
            instruction_style
        ))

        elements.append(Paragraph(
            "<b>Answer All Questions from Part-A and Part-B</b>",
            instruction_style
        ))

        elements.append(Spacer(1, 6))

        # -------- PART TITLE STYLE --------
        part_title_style = ParagraphStyle(
            name="PartTitle",
            parent=styles["Normal"],
            fontSize=12,
            leading=14,
            alignment=TA_CENTER,
            spaceAfter=6
        )

        # -------- QUESTION ALIGNMENT STYLES --------
        normal_style = styles["Normal"]

        right_align_style = ParagraphStyle(
            name="RightAlign",
            parent=styles["Normal"],
            alignment=TA_RIGHT
        )

        # Proper hanging indent like sample paper
        question_style = ParagraphStyle(
            name="QuestionStyle",
            parent=styles["Normal"],
            fontSize=11,
            leading=14,
            leftIndent=25,
            firstLineIndent=-25,
        )

        # -------- QUESTIONS --------

        # Convert (2M) → [2]
        content = re.sub(r"\((\d+)M\)", r"[\1]", content)

        # First roman question keeps 1.
        content = re.sub(r"1i\)", "1. i.", content)

        # Remaining roman numerals remove the 1
        content = re.sub(r"1(ii|iii|iv|v)\)", r"\1.", content)

        # Convert 2a) → 2. (a)
        content = re.sub(r"(\d+)([a-z])\)", r"\1. (\2)", content)

        for line in content.split("\n"):

            clean_line = line.strip()

            if not clean_line:
                elements.append(Spacer(1, 3))
                continue

            if clean_line.upper().startswith("PART-A"):
                elements.append(
                    Paragraph("<b>Part-A (Max Marks: 10)</b>", part_title_style)
                )
                elements.append(Spacer(1, 4))
                continue

            if clean_line.upper().startswith("PART-B"):
                elements.append(
                    Paragraph("<b>Part-B (Max Marks: 30)</b>", part_title_style)
                )
                elements.append(Spacer(1, 4))
                continue

            if clean_line.upper() == "OR":
                elements.append(Spacer(1, 4))
                elements.append(
                    Paragraph("<b>OR</b>", part_title_style)
                )
                elements.append(Spacer(1, 4))
                continue

            # ---- FIND MARKS LIKE [2] OR [5] ----
            match = re.search(r"\[\d+\]", clean_line)

            if match:

                marks = match.group()
                question_text = clean_line.replace(marks, "").strip()

                question_table = Table(
                    [
                        [
                            Paragraph(question_text, question_style),
                            Paragraph(marks, right_align_style)
                        ]
                    ],
                    colWidths=[6.1 * inch, 0.4 * inch]
                )

                question_table.setStyle(
                    TableStyle([
                        ("VALIGN", (0,0), (-1,-1), "TOP"),
                        ("ALIGN", (1,0), (1,0), "RIGHT"),

                        ("LEFTPADDING", (0,0), (-1,-1), 0),
                        ("RIGHTPADDING", (0,0), (-1,-1), 0),

                        ("TOPPADDING", (0,0), (-1,-1), 0),
                        ("BOTTOMPADDING", (0,0), (-1,-1), 0),
                    ])
                )

                elements.append(question_table)
                elements.append(Spacer(1, 5))

            else:
                elements.append(
                    Paragraph(clean_line, normal_style)
                )

        doc.build(elements)
        buffer.seek(0)   

        return send_file(
            buffer,
            as_attachment=True,
            download_name=f"Question_Paper_{paper_id}.pdf",
            mimetype="application/pdf"
        )
    
    # =====================================================
    # ====================== DOCX =========================
    # =====================================================
    elif format_type == "docx":
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from io import BytesIO
        import os

        buffer = BytesIO()
        document = Document()

        # ---------- PAGE MARGINS ----------
        section = document.sections[0]
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

                # ---------- LOGO (CENTERED, NO TABLE) ----------
        logo_para = document.add_paragraph()
        logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_run = logo_para.add_run()
        logo_run.add_picture("static/logo.png", width=Inches(5.5))

        # Remove spacing below logo
        logo_para.paragraph_format.space_after = 0
        logo_para.paragraph_format.space_before = 0

                # ---------- REGULATION BOX (MATCH SAMPLE EXACTLY) ----------
        reg_table = document.add_table(rows=1, cols=1)
        reg_table.autofit = False
        reg_table.alignment = WD_TABLE_ALIGNMENT.RIGHT
        reg_table.columns[0].width = Inches(0.8)   # smaller width like sample

        reg_cell = reg_table.cell(0, 0)
        reg_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        reg_para = reg_cell.paragraphs[0]
        reg_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Remove internal spacing
        reg_para.paragraph_format.space_before = 0
        reg_para.paragraph_format.space_after = 0

        reg_run = reg_para.add_run(regulation)
        reg_run.bold = True
        reg_run.font.size = Pt(11)   # slightly compact

        # Tight compact border
        tc = reg_cell._tc
        tcPr = tc.get_or_add_tcPr()

        # Reduce cell margins (VERY IMPORTANT)
        tcMar = OxmlElement('w:tcMar')
        for m in ('top', 'left', 'bottom', 'right'):
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), "60")   # small padding
            node.set(qn('w:type'), "dxa")
            tcMar.append(node)
        tcPr.append(tcMar)

        borders = OxmlElement('w:tcBorders')

        for border_name in ('top', 'left', 'bottom', 'right'):
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '20')   # thicker border like sample
            border.set(qn('w:color'), '000000')
            borders.append(border)

        tcPr.append(borders)
        document.add_paragraph("")
    
        # ---------- STREAM ----------
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{stream}; {semester}; {exam_type}, {month_year}")
        run.bold = True

        # ---------- SUBJECT ----------
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(subject.upper())
        run.bold = True
        run.font.size = Pt(14)

        document.add_paragraph("")

        # ---------- TIME / MARKS TABLE ----------
        info_table = document.add_table(rows=2, cols=2)
        info_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        info_table.autofit = False

        info_table.columns[0].width = Inches(3.25)
        info_table.columns[1].width = Inches(3.25)

        info_table.cell(0, 0).text = f"Time: {time_duration}"
        info_table.cell(0, 1).text = f"Max. Marks: {max_marks}"
        info_table.cell(1, 0).text = f"Branch & Section: {branch_section}"
        info_table.cell(1, 1).text = f"Date: {exam_date}"

        # Remove borders
        for row in info_table.rows:
            for cell in row.cells:
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                tcBorders = OxmlElement('w:tcBorders')
                for border_name in ('top', 'left', 'bottom', 'right'):
                    border = OxmlElement(f'w:{border_name}')
                    border.set(qn('w:val'), 'nil')
                    tcBorders.append(border)
                tcPr.append(tcBorders)

        document.add_paragraph("")

        # ---------- REGISTER NUMBER ----------
        register_table = document.add_table(rows=1, cols=2)
        register_table.alignment = WD_TABLE_ALIGNMENT.LEFT
        register_table.autofit = False
        register_table.columns[0].width = Inches(1.6)
        register_table.columns[1].width = Inches(4.8)

        label_cell = register_table.cell(0, 0)
        label_para = label_cell.paragraphs[0]
        label_run = label_para.add_run("Registered No")
        label_run.bold = True
        label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        boxes_cell = register_table.cell(0, 1)
        boxes_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        boxes_table = boxes_cell.add_table(rows=1, cols=8)
        boxes_table.autofit = False

        for i in range(8):
            boxes_table.columns[i].width = Inches(0.4)

        for cell in boxes_table.rows[0].cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for border_name in ('top', 'left', 'bottom', 'right'):
                border = OxmlElement(f'w:{border_name}')
                border.set(qn('w:val'), 'single')
                border.set(qn('w:sz'), '12')
                border.set(qn('w:color'), '000000')
                borders.append(border)
            tcPr.append(borders)

        # Remove outer borders
        for row in register_table.rows:
            for cell in row.cells:
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                tcBorders = OxmlElement('w:tcBorders')
                for border_name in ('top', 'left', 'bottom', 'right'):
                    border = OxmlElement(f'w:{border_name}')
                    border.set(qn('w:val'), 'nil')
                    tcBorders.append(border)
                tcPr.append(tcBorders)

        document.add_paragraph("")
        
        # ---------- HORIZONTAL LINE ----------
        line_para = document.add_paragraph()
        p = line_para._p
        pPr = p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '12')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '000000')
        pBdr.append(bottom)
        pPr.append(pBdr)

        document.add_paragraph("")

        # ---------- INSTRUCTIONS ----------
        p1 = document.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run1 = p1.add_run("Question Paper Consists of Part-A and Part-B")
        run1.bold = True

        p2 = document.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run("Answer All Questions from Part-A and Part-B")
        run2.bold = True

        document.add_paragraph("")

        # ---------- QUESTIONS ----------
        for line in content.split("\n"):
            clean_line = line.strip()

            if not clean_line:
                continue

            if clean_line.upper().startswith("PART-A"):
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run("Part-A (Max Marks: 10)")
                run.bold = True
                document.add_paragraph("")
                continue

            if clean_line.upper().startswith("PART-B"):
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run("Part-B (Max Marks: 30)")
                run.bold = True
                document.add_paragraph("")
                continue
            
            if clean_line.upper() == "OR":
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run("OR")
                run.bold = True
                continue

            document.add_paragraph(clean_line)

        document.save(buffer)
        buffer.seek(0)

        return send_file(
            buffer,
            as_attachment=True,
            download_name=f"Question_Paper_{paper_id}.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    else:
        return {"error": "Format must be pdf or docx"}, 400

@app.route("/questions", methods=["GET"])
def get_questions():
    questions = Question.query.all()

    result = []
    for q in questions:
        result.append({
            "id": q.id,
            "unit": q.unit,
            "marks": q.marks,
            "text": q.question_text
        })

    return {"questions": result}

from flask import request

@app.route("/add-question", methods=["POST"])
def add_question():

    data = request.json

    unit = data.get("unit")
    marks = data.get("marks")
    text = data.get("text")

    if not unit or not marks or not text:
        return {"error": "Missing fields"}, 400

    new_question = Question(
        unit=unit,
        marks=marks,
        question_text=text
    )

    db.session.add(new_question)
    db.session.commit()

    return {"message": "Question added successfully"}

@app.route("/delete-question/<int:question_id>", methods=["DELETE"])
def delete_question(question_id):

    question = Question.query.get(question_id)   # ✅ FIXED

    if not question:
        return {"error": "Question not found"}, 404

    db.session.delete(question)
    db.session.commit()

    return {"message": "Question deleted successfully"}

@app.route("/update-question/<int:question_id>", methods=["PUT"])
def update_question(question_id):

    question = Question.query.get(question_id)

    if not question:
        return {"error": "Question not found"}, 404

    data = request.get_json()

    question.unit = data.get("unit", question.unit)
    question.marks = data.get("marks", question.marks)
    question.question_text = data.get("text", question.question_text) 

    db.session.commit()

    return {"message": "Question updated successfully"}

# ------------------ MAIN ------------------

if __name__ == "__main__":
    with app.app_context():
        db.create_all()
    app.run(debug=True)